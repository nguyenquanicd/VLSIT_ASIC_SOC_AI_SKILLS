#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# File     : vlsit_csr_tool.py
# Function : fill, validate and dump the VLSIT CSR specification workbook from a JSON register description
# Author   : VLSI Technology
# GitHub   : https://github.com/nguyenquanicd/VLSIT_CODEX_CLAUDE_SKILLS
# ---------------------------------------------------------------------------
"""Fill, check and summarize the VLSIT CSR specification workbook from a JSON register description.

Purpose : the register map lives only in the workbook VLSIT_CSR_Specification_Template.xlsx (sheets CSR,
          ASYNC_CSR, Register_Description). This tool writes it from a JSON file, validates the map, and
          prints a compact Markdown summary so the assistant does not have to read the workbook cells.
Actions :
  validate REGS.json            check the JSON description (offsets, bit fields, types, resets)
  fill REGS.json --out FILE.xlsx [--template FILE.xlsx] [--auto-reserved] [--force]
                                write the workbook from the template (sheet CSR, or ASYNC_CSR when async=1)
  dump FILE.xlsx [--sheet CSR]  print the register map of a workbook as Markdown
  template                      print a starter JSON description
  ports REGS.json               print the ports of the module that APB-CSR-Generator builds from the
                                description (APB ports, one set of ports per field) as Markdown
Input   : JSON {"module": "name", "async": 0, "registers": [{"name": "ctrl", "offset": "0x0",
          "description": "...", "fields": [{"name": "en", "bits": "[0]", "type": "rw", "reset": "0x0",
          "description": "..."}]}]}. Protocol APB, data width 32, address width 16, write strobe 4 are fixed
          by the workbook (cells B4:B7).
Rules   : types rw, ro, rwi, w1c; offsets unique, 4-byte aligned, below 2^16; fields of one register tile
          bits 31..0 without gap or overlap (--auto-reserved adds 'reserved' ro fields for gaps); reset
          fits the field width; register and field names are lower-case identifiers (they become port and
          localparam names of APB-CSR-Generator); a field named 'reserved' is type ro; a generated port name
          longer than 30 characters, or a name that spells out a NAM-08 word (status, enable ...), is a warning.
Exit    : 0 success, 1 findings or error, 3 openpyxl missing (python -m pip install openpyxl).
Text files are read and written as UTF-8.
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "reference" / "templates" / "VLSIT_CSR_Specification_Template.xlsx"
TYPES = {"rw", "ro", "rwi", "w1c"}
NAME_RE = re.compile(r"[a-z][a-z0-9_]*")
MAX_NAME = 30             # NAM-07 of the RTL design rule
NAM08 = {"address": "addr", "next": "nxt", "counter": "cnt", "enable": "en", "clock": "clk", "request": "req",
         "response": "resp", "select": "sel", "control": "ctrl", "status": "sts", "error": "err", "frequency": "freq",
         "divider": "div", "threshold": "thr", "arbiter": "arb", "grant": "gnt", "priority": "prio", "multiplexer": "mux",
         "decoder": "dec", "encoder": "enc", "maximum": "max", "minimum": "min", "number": "num", "parity": "par",
         "configuration": "config", "configure": "config", "interrupt": "int", "reset": "rst", "acknowledge": "ack",
         "command": "cmd", "buffer": "buf", "pointer": "ptr", "index": "idx", "length": "len", "source": "src",
         "destination": "dst", "transmit": "tx", "receive": "rx", "write": "wr", "read": "rd"}   # NAM-08 of the RTL rule
FIRST_ROW = 16            # first register row of the template
SAMPLE = {"module": "m_vlsi_csr", "async": 0, "registers": [
    {"name": "ctrl", "offset": "0x0", "description": "Control", "fields": [
        {"name": "enable", "bits": "[0]", "type": "rw", "reset": "0x0", "description": "Enable"},
        {"name": "reserved", "bits": "[31:1]", "type": "ro", "reset": "0x0", "description": ""}]}]}


def parse_bits(text):
    m = re.fullmatch(r"\[(\d+)(?::(\d+))?\]", text.strip())
    if not m:
        raise ValueError("bits '%s' must look like [7] or [7:4]" % text)
    hi = int(m.group(1))
    lo = int(m.group(2)) if m.group(2) is not None else hi
    if hi < lo or hi > 31:
        raise ValueError("bits '%s' must satisfy 31 >= high >= low >= 0" % text)
    return hi, lo


def to_int(text):
    return int(str(text), 0)


def check(doc, auto_reserved=False):
    errors, seen_names, offsets = [], set(), {}
    if not re.fullmatch(r"[A-Za-z_]\w*", str(doc.get("module", ""))):
        errors.append("module name must be an identifier")
    if doc.get("async", 0) not in (0, 1):
        errors.append("async must be 0 or 1")
    for reg in doc.get("registers", []):
        name = reg.get("name", "?")
        if not NAME_RE.fullmatch(str(name)):
            errors.append("register name %s must be a lower-case identifier (letters, digits, underscore)" % name)
        if name in seen_names:
            errors.append("duplicate register name %s" % name)
        seen_names.add(name)
        try:
            off = to_int(reg["offset"])
            if off % 4 or not 0 <= off < 1 << 16:
                errors.append("%s: offset %s must be 4-byte aligned and below 0x10000" % (name, reg["offset"]))
            if off in offsets:
                errors.append("%s: offset %s already used by %s" % (name, reg["offset"], offsets[off]))
            offsets[off] = name
        except (KeyError, ValueError):
            errors.append("%s: missing or bad offset" % name)
        covered, fnames = {}, set()
        for f in reg.get("fields", []):
            fname = f.get("name", "?")
            if fname in fnames and fname != "reserved":
                errors.append("%s: duplicate field %s" % (name, fname))
            fnames.add(fname)
            if not NAME_RE.fullmatch(str(fname)):
                errors.append("%s.%s: field name must be a lower-case identifier" % (name, fname))
            if fname == "reserved" and f.get("type") != "ro":
                errors.append("%s.reserved: a reserved field must be type ro (the generator would add ports to another type)" % name)
            if f.get("type") not in TYPES:
                errors.append("%s.%s: type must be one of %s" % (name, fname, sorted(TYPES)))
            try:
                hi, lo = parse_bits(f["bits"])
            except (KeyError, ValueError) as exc:
                errors.append("%s.%s: %s" % (name, fname, exc))
                continue
            for b in range(lo, hi + 1):
                if b in covered:
                    errors.append("%s.%s: bit %d also in %s" % (name, fname, b, covered[b]))
                covered[b] = fname
            try:
                if to_int(f.get("reset", "0x0")) >= 1 << (hi - lo + 1):
                    errors.append("%s.%s: reset %s does not fit %d bit(s)" % (name, fname, f.get("reset"), hi - lo + 1))
            except ValueError:
                errors.append("%s.%s: bad reset value" % (name, fname))
        gaps = [b for b in range(32) if b not in covered]
        if gaps and not auto_reserved:
            errors.append("%s: bits %s are not covered; add a reserved field or use --auto-reserved" % (name, gaps))
    return errors


def with_reserved(reg):
    fields = [dict(f) for f in reg["fields"]]
    covered = set()
    for f in fields:
        hi, lo = parse_bits(f["bits"])
        covered.update(range(lo, hi + 1))
    b = 31
    while b >= 0:
        if b in covered:
            b -= 1
            continue
        hi = b
        while b >= 0 and b not in covered:
            b -= 1
        fields.append({"name": "reserved", "bits": "[%d:%d]" % (hi, b + 1) if hi != b + 1 else "[%d]" % hi,
                       "type": "ro", "reset": "0x0", "description": ""})
    fields.sort(key=lambda f: -parse_bits(f["bits"])[0])
    return fields


def register_reset(fields):
    value = 0
    for f in fields:
        hi, lo = parse_bits(f["bits"])
        value |= (to_int(f.get("reset", "0x0")) & ((1 << (hi - lo + 1)) - 1)) << lo
    return "0x%08X" % value


def field_ports(reg, f):
    """Ports that APB-CSR-Generator builds for one field: list of (name, direction, width)."""
    hi, lo = parse_bits(f["bits"])
    width, base = hi - lo + 1, "%s_%s" % (reg["name"], f["name"])
    if f["name"] == "reserved":
        return []
    if f["type"] == "ro":
        return [("i_" + base, "input", width)]
    ports = [("o_" + base, "output", width)]
    if f["type"] in ("rwi", "w1c"):
        ports += [("i_hw_wdata_" + base, "input", width), ("i_hw_we_" + base, "input", 1)]
    return ports


def name_warnings(doc):
    warns = []
    for reg in doc.get("registers", []):
        labels = [("register", reg.get("name", ""))] + [("field", "%s.%s" % (reg.get("name", ""), f.get("name", ""))) for f in reg.get("fields", [])]
        for label, name in labels:
            for word in str(name).replace(".", "_").split("_"):
                if word in NAM08:
                    warns.append("%s %s spells out '%s'; NAM-08 of the RTL rule writes '%s'" % (label, name, word, NAM08[word]))
    for reg in doc.get("registers", []):
        longest = ["LPARA_%s_OFFSET" % reg["name"].upper(), "p_com_%s_value" % reg["name"], "w_we_%s" % reg["name"]]
        for f in reg.get("fields", []):
            try:
                longest += [p[0] for p in field_ports(reg, f)]
                if f["name"] != "reserved":
                    longest += ["reg_%s_%s" % (reg["name"], f["name"]), "w_%s_%s_nxt" % (reg["name"], f["name"])]
            except (KeyError, ValueError):
                pass
        for n in longest:
            if len(n) > MAX_NAME:
                warns.append("%s has %d characters (limit %d); shorten the register or field name" % (n, len(n), MAX_NAME))
    return warns


def need_openpyxl():
    if load_workbook is None:
        print("ERROR: openpyxl missing; install with: python -m pip install openpyxl (ask the user first)")
        return False
    return True


def action_validate(a):
    doc = json.loads(Path(a.json).read_text(encoding="utf-8"))
    errors = check(doc, a.auto_reserved)
    for e in errors:
        print("ERROR: " + e)
    warns = name_warnings(doc) if not errors else []
    for w in warns:
        print("WARN: " + w)
    print("%s: %d error(s), %d warning(s), %d register(s)" % ("PASS" if not errors else "FAIL", len(errors), len(warns), len(doc.get("registers", []))))
    return 1 if errors else 0


def copy_style(src, dst):
    dst._style = copy.copy(src._style)


def action_fill(a):
    if not need_openpyxl():
        return 3
    doc = json.loads(Path(a.json).read_text(encoding="utf-8"))
    errors = check(doc, a.auto_reserved)
    if errors:
        for e in errors:
            print("ERROR: " + e)
        print("FAIL: fix the register description first")
        return 1
    out = Path(a.out)
    if out.exists() and not a.force:
        print("ERROR: %s already exists; pass --force to overwrite it" % out)
        return 1
    wb = load_workbook(a.template or TEMPLATE)
    ws = wb["ASYNC_CSR" if doc.get("async", 0) == 1 else "CSR"]
    reg_style = [ws.cell(FIRST_ROW, c) for c in range(1, 8)]
    fld_style = [ws.cell(FIRST_ROW + 1, c) for c in range(1, 8)]
    reg_style, fld_style = [copy.copy(c._style) for c in reg_style], [copy.copy(c._style) for c in fld_style]
    for row in range(FIRST_ROW, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(row, c).value = None
    ws["B3"], ws["B8"] = doc["module"], str(doc.get("async", 0))
    row = FIRST_ROW
    for reg in sorted(doc["registers"], key=lambda r: to_int(r["offset"])):
        fields = with_reserved(reg) if a.auto_reserved else sorted(reg["fields"], key=lambda f: -parse_bits(f["bits"])[0])
        values = [reg["name"], "0x%X" % to_int(reg["offset"]), None, None, None, register_reset(fields), reg.get("description", "")]
        for c, v in enumerate(values, 1):
            ws.cell(row, c).value, ws.cell(row, c)._style = v, copy.copy(reg_style[c - 1])
        row += 1
        for f in fields:
            vals = [None, None, f["name"], f["bits"], f["type"], f.get("reset", "0x0"), f.get("description", "")]
            for c, v in enumerate(vals, 1):
                ws.cell(row, c).value, ws.cell(row, c)._style = v, copy.copy(fld_style[c - 1])
            row += 1
    if ws.max_row >= row:
        # the generator reads every row up to max_row; styled empty rows of the template sample would crash it
        ws.delete_rows(row, ws.max_row - row + 1)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("OK: wrote %s (sheet %s, %d registers, %d rows)" % (out, ws.title, len(doc["registers"]), row - FIRST_ROW))
    return 0


def action_dump(a):
    if not need_openpyxl():
        return 3
    ws = load_workbook(a.xlsx, data_only=True)[a.sheet]
    cfg = {str(ws.cell(r, 1).value): ws.cell(r, 2).value for r in range(3, 9)}
    print("Module: %s; protocol %s; data width %s; address width %s; asynchronous %s\n" % (
        cfg.get("Module Name"), cfg.get("Protocol"), cfg.get("Data Width"), cfg.get("Address Width"), cfg.get("Asynchronous")))
    print("| Register | Offset | Field | Bits | Type | Reset | Description |\n|---|---|---|---|---|---|---|")
    for r in range(FIRST_ROW, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if any(v is not None for v in vals):
            print("| " + " | ".join("" if v is None else str(v).replace("\n", " ").replace("|", "\\|") for v in vals) + " |")
    return 0


def action_ports(a):
    doc = json.loads(Path(a.json).read_text(encoding="utf-8"))
    errors = check(doc, True)
    if errors:
        for e in errors:
            print("ERROR: " + e)
        print("FAIL: fix the register description first")
        return 1
    asyn = doc.get("async", 0) == 1
    regclk = "i_clk_reg" if asyn else "i_clk_apb"
    rows = [("i_clk_apb", "input", 1, "APB clock", "-"), ("i_rst_n_apb", "input", 1, "APB reset, active low", "-")]
    if asyn:
        rows += [("i_clk_reg", "input", 1, "Register clock", "-"), ("i_rst_n_reg", "input", 1, "Register reset, active low", "-")]
    rows += [("i_apb_paddr", "input", 16, "APB address", "-"), ("i_apb_pwdata", "input", 32, "APB write data", "-"),
             ("i_apb_pwrite", "input", 1, "APB write", "-"), ("i_apb_penable", "input", 1, "APB enable", "-"),
             ("i_apb_psel", "input", 1, "APB select", "-"), ("i_apb_pstrb", "input", 4, "APB byte strobes", "-"),
             ("i_apb_pprot", "input", 3, "APB protection (bit 1 used)", "-"),
             ("i_protect_en", "input", 1, "1 = non-secure access gets PSLVERR", "-"),
             ("i_slverr_en", "input", 1, "1 = PSLVERR checks on", "-"),
             ("o_apb_pready", "output", 1, "APB ready", "-"), ("o_apb_pslverr", "output", 1, "APB error", "-"),
             ("o_apb_prdata", "output", 32, "APB read data", "-")]
    kind = {"ro": "status in", "rw": "control out", "rwi": "control out, hardware load", "w1c": "status/event, hardware load"}
    for reg in sorted(doc["registers"], key=lambda r: to_int(r["offset"])):
        for f in reg["fields"]:
            for n, d, w in field_ports(reg, f):
                role = "hardware value" if n.startswith("i_hw_wdata") else "hardware load enable" if n.startswith("i_hw_we") else kind[f["type"]]
                rows.append((n, d, w, "%s.%s (%s)" % (reg["name"], f["name"], role), regclk))
    print("Module %s (%s); register clock %s\n" % (doc["module"], "asynchronous" if asyn else "single clock", regclk))
    print("| Port | Direction | Width | Meaning | Clock domain of the signal |\n|---|---|---:|---|---|")
    for n, d, w, m, c in rows:
        print("| %s | %s | %d | %s | %s |" % (n, d, w, m, c if c != "-" else "i_clk_reg" if n in ("i_clk_reg", "i_rst_n_reg") else "i_clk_apb"))
    return 0


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="action", required=True)
    v = sub.add_parser("validate")
    v.add_argument("json")
    v.add_argument("--auto-reserved", action="store_true")
    f = sub.add_parser("fill")
    f.add_argument("json")
    f.add_argument("--out", required=True)
    f.add_argument("--template")
    f.add_argument("--auto-reserved", action="store_true")
    f.add_argument("--force", action="store_true")
    d = sub.add_parser("dump")
    d.add_argument("xlsx")
    d.add_argument("--sheet", default="CSR")
    sub.add_parser("template")
    pt = sub.add_parser("ports")
    pt.add_argument("json")
    a = parser.parse_args()
    if a.action == "template":
        print(json.dumps(SAMPLE, indent=2))
        return 0
    return {"validate": action_validate, "fill": action_fill, "dump": action_dump, "ports": action_ports}[a.action](a)


if __name__ == "__main__":
    sys.exit(main())
